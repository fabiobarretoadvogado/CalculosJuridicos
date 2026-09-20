import copy
import csv
import io
import json
from datetime import date
from decimal import Decimal as D
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient
from pypdf import PdfReader

from liquidacao_custom.api.main import app
from liquidacao_custom.core import motor_ipca_taxa_legal as legal
from liquidacao_custom.core.criterios_simplificados import CalculoSimplificado
from liquidacao_custom.core.exportadores import exportar_csv, exportar_excel
from liquidacao_custom.core.motor_simplificado import executar_calculo, moeda, criterios_publicos
from liquidacao_custom.core.relatorio_pdf import exportar_pdf

CASO = Path(__file__).resolve().parents[1] / "examples" / "processo_3005719-41.2026.8.06.0297.json"


def entrada():
    return json.loads(CASO.read_text(encoding="utf-8"))


def calcular(payload):
    return executar_calculo(CalculoSimplificado.model_validate(payload))


def test_caso_real_preserva_datas_e_reproduz_taxa_do_bcb():
    r = calcular(entrada())
    p = r.parcelas[0]
    assert p.componentes["taxa_legal"].data_inicial == date(2026, 2, 2)
    assert p.componentes["ipca"].data_inicial == date(2026, 8, 28)
    assert p.componentes["taxa_legal"].taxa_acumulada_percentual == D("4.325078")
    assert p.componentes["ipca"].data_final == date(2026, 8, 30)
    assert p.componentes["taxa_legal"].base_calculo == D("5000") * D("0.9968") ** (D(3)/31)
    assert p.correcao_monetaria == D("-1.55")
    assert p.juros_mora == D("216.19")
    assert r.resumo.total_atualizado == D("5214.64")
    assert set(p.componentes) == {"ipca", "taxa_legal"}
    assert r.resumo.multas == r.resumo.honorarios == r.resumo.custas == 0
    assert sum((m.juros_periodo for m in p.memoria_juros), D(0)) == p.juros_mora
    assert len(p.memoria_juros) == 7 and len(p.memoria_correcao) == 1
    assert r.premissas["sha256_ipca"] and r.premissas["sha256_taxa_legal"]


@pytest.mark.parametrize("criterio", ["citacao", "data_fixa", "por_parcela"])
def test_inicio_anterior_nao_e_cortado_no_novo_perfil(criterio):
    payload = entrada()
    payload["dados_gerais"]["criterio_inicio_juros"] = criterio
    payload["parcelas"][0]["data_inicial_juros"] = "2026-02-02"
    calculo = CalculoSimplificado.model_validate(payload)
    assert calculo.inicio_juros_parcela(calculo.parcelas[0]) == date(2026, 2, 2)
    payload["perfil"] = "selic_ipcae_poupanca_v1"
    antigo = CalculoSimplificado.model_validate(payload)
    assert antigo.inicio_juros_parcela(antigo.parcelas[0]) == date(2026, 8, 28)


def test_confronto_exemplo_oficial_um_dia_sem_dia_a_mais():
    payload = entrada()
    payload["dados_gerais"].update(data_base="2024-08-31", data_inicial_juros="2024-08-30")
    payload["parcelas"][0].update(data_vencimento="2024-08-31", valor_bruto="1000")
    p = calcular(payload).parcelas[0]
    assert p.componentes["taxa_legal"].taxa_acumulada_percentual == D("0.019526")
    assert p.juros_mora == D("0.20")
    assert p.correcao_monetaria == 0
    assert p.componentes["ipca"].data_inicial is None


def test_mesmo_dia_nao_gera_encargos():
    payload = entrada()
    payload["dados_gerais"]["data_inicial_juros"] = "2026-08-31"
    payload["parcelas"][0]["data_vencimento"] = "2026-08-31"
    r = calcular(payload)
    assert r.resumo.total_atualizado == 5000
    assert not r.memoria_mensal
    assert all(c.data_inicial is None for c in r.parcelas[0].componentes.values())


def test_juros_posteriores_a_correcao():
    payload = entrada()
    payload["parcelas"][0]["data_vencimento"] = "2026-02-02"
    payload["dados_gerais"]["data_inicial_juros"] = "2026-08-28"
    p = calcular(payload).parcelas[0]
    assert len(p.memoria_correcao) == 7 and len(p.memoria_juros) == 1
    assert p.componentes["taxa_legal"].taxa_acumulada_percentual == D("0.111728")


def test_pagamento_integral_nao_gera_divida():
    payload = entrada()
    payload["parcelas"][0]["valor_pago_na_data"] = "5000"
    r = calcular(payload)
    assert r.resumo.total_atualizado == 0
    assert r.parcelas[0].componentes["taxa_legal"].valor == 0


def test_taxa_zero_oficial_nao_vira_um_por_cento():
    payload = entrada()
    payload["dados_gerais"].update(data_base="2025-03-31", data_inicial_juros="2025-03-01")
    payload["parcelas"][0]["data_vencimento"] = "2025-03-31"
    p = calcular(payload).parcelas[0]
    assert p.componentes["taxa_legal"].data_inicial is not None
    assert p.juros_mora == 0
    assert p.memoria_juros and "0.000000%" in p.memoria_juros[0].observacao


@pytest.mark.parametrize("data", ["2024-08-29", "2023-01-01"])
def test_nao_estende_taxa_legal_para_periodo_anterior(data):
    payload = entrada()
    payload["dados_gerais"]["data_inicial_juros"] = data
    with pytest.raises(ValueError, match="30/08/2024"):
        calcular(payload)


def test_indices_ausentes_e_negativos_impedem_substituicao(tmp_path, monkeypatch):
    arquivo = tmp_path / "taxas.json"
    dados = json.loads(legal.BASE_TAXA_LEGAL.read_text(encoding="utf-8"))
    monkeypatch.setattr(legal, "BASE_TAXA_LEGAL", arquivo)
    del dados["taxa_legal"]["2026-04"]
    arquivo.write_text(json.dumps(dados), encoding="utf-8")
    with pytest.raises(ValueError, match="Índice ausente"):
        calcular(entrada())
    dados["taxa_legal"]["2026-04"] = "-1"
    arquivo.write_text(json.dumps(dados), encoding="utf-8")
    with pytest.raises(ValueError, match="negativa"):
        calcular(entrada())


def test_ipca_futuro_nao_e_estimado():
    payload = entrada()
    payload["dados_gerais"]["data_base"] = "2026-09-17"
    with pytest.raises(ValueError, match="último mês com IPCA"):
        calcular(payload)
    assert criterios_publicos(perfil="ipca_taxa_legal_v1")["data_base_maxima"] == "2026-08-31"


def test_desconto_multa_honorarios_e_custas_usam_ordem_existente():
    payload = entrada()
    payload["parcelas"][0]["multa_percentual"] = "2"
    payload["descontos"] = {"itens": [{"numero": 1, "descricao": "Teste", "data": "2026-08-29", "valor": "100"}]}
    payload["custas_despesas"] = [{"numero": 1, "nome": "Teste", "data": "2026-08-29", "valor": "300"}]
    payload["honorarios_sucumbenciais"] = {"aplicar": True, "base": "proveito_economico", "percentual": "20"}
    r = calcular(payload)
    credito = r.parcelas[0].total_parcela - r.resumo.abatimentos
    assert r.honorarios_sucumbenciais.base_atualizada == credito
    assert r.resumo.total_atualizado == credito + r.honorarios_sucumbenciais.valor + r.resumo.custas
    assert r.parcelas[0].multa_detalhes.juros_mora > 0
    assert r.descontos.premissas["perfil"] == "ipca_taxa_legal_v1"


def test_api_pdf_excel_csv_mesmos_valores_e_fontes(tmp_path):
    r = calcular(entrada())
    c = TestClient(app)
    response = c.post("/api/v1/calculo", json=entrada())
    assert response.status_code == 200
    assert response.json()["resumo"]["total_atualizado"] == "5214.64"
    pdf = c.post("/api/v1/calculo/exportar/pdf", json=entrada())
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    text = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf.content)).pages)
    assert "4,325078%" in text and "Memória da Taxa Legal" in text
    assert "02/02/2026" in text and "28/08/2026" in text
    assert "Poupança" not in text and "IPCA-E pós EC" not in text
    assert "5.214,64" in text and "BCB" in text
    xlsx = tmp_path / "memoria.xlsx"
    exportar_excel(r, str(xlsx))
    wb = openpyxl.load_workbook(xlsx)
    assert "Taxa Legal" in wb.sheetnames and wb["Taxa Legal"].max_row == 8
    assert wb["Taxa Legal"]["E2"].value == 0.962232
    assert wb["Taxa Legal"]["E2"].number_format == "0.000000"
    assert wb["Índices por parcela"]["H3"].number_format == "0.000000"
    assert wb["Parcelas"]["B1"].value == "Correção desde"
    wb.close()
    exportar_csv(r, str(tmp_path / "csv"))
    rows = list(csv.reader((tmp_path / "csv" / "taxa_legal_mensal.csv").open(encoding="utf-8-sig")))
    assert rows[1][4] == "0.962232"
    assert D(rows[-1][6]) == D("4.325078")


@pytest.mark.parametrize("quantidade", [20, 36])
def test_pdf_muitas_parcelas_preserva_cabecalhos_e_datas(quantidade):
    payload = entrada()
    payload["dados_gerais"]["observacoes"] = "Observação extensa de conferência. " * 20
    payload["parcelas"] = [dict(payload["parcelas"][0], numero=n, historico=f"Dano moral {n}") for n in range(1, quantidade+1)]
    r = calcular(payload)
    doc = PdfReader(io.BytesIO(exportar_pdf(r)))
    text = "\n".join(p.extract_text() for p in doc.pages)
    assert text.count("Correção desde") == 1
    assert text.count("Memória da Taxa Legal - parcela") == quantidade
    assert len(doc.pages) >= 2
    assert r.resumo.total_atualizado == D("5214.64") * quantidade


def test_pdf_cabecalho_nao_fica_orfao_apos_observacao_extensa():
    payload = entrada()
    payload["dados_gerais"]["observacoes"] = "Observação extensa para conferir margens, paginação e fontes. " * 30
    payload["parcelas"] = [dict(payload["parcelas"][0], numero=n, historico=f"Dano moral {n}") for n in range(1, 21)]
    paginas = [" ".join(p.extract_text().split()) for p in PdfReader(io.BytesIO(exportar_pdf(calcular(payload)))).pages]
    inicio = next(t for t in paginas if "Correção desde" in t)
    assert "Dano moral 1" in inicio
    assert "Valores por parcela" in inicio
