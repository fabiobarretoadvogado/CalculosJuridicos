import copy
import io
import json
import zipfile
from decimal import Decimal as D

import openpyxl
import pytest
from fastapi.testclient import TestClient
from pydantic import ValidationError
from pypdf import PdfReader

from liquidacao_custom.api.main import app
from liquidacao_custom.core.criterios_simplificados import CalculoSimplificado
from liquidacao_custom.core.motor_simplificado import executar_calculo
from liquidacao_custom.core.exportadores import exportar_excel


def entrada(perfil='selic_ipcae_poupanca_v1', perfil_descontos=None):
    return {
        'perfil': perfil,
        'dados_gerais': {'data_base': '2026-08-31', 'criterio_inicio_juros': 'vencimento'},
        'parcelas': [{'numero': 1, 'historico': 'Parcela original', 'data_vencimento': '2025-09-01', 'valor_bruto': '1000', 'valor_pago_na_data': '200'}],
        'descontos': {'perfil': perfil_descontos, 'itens': [
            {'numero': 1, 'descricao': 'Pagamento posterior', 'data': '2025-10-15', 'valor': '150'},
            {'numero': 2, 'descricao': 'Crédito posterior', 'data': '2026-02-10', 'valor': '50'},
        ]},
        'custas_despesas': [{'numero': 1, 'nome': 'Taxa', 'data': '2025-10-01', 'valor': '100'}],
    }


def calcular(payload):
    return executar_calculo(CalculoSimplificado.model_validate(payload))


@pytest.mark.parametrize('perfil', ['selic_ipcae_poupanca_v1', 'selic_ipcae_2aa_v1', 'selic_cjf_v1', 'ipcae_1am_simples_v1'])
@pytest.mark.parametrize('perfil_descontos', [None, 'selic_cjf_v1'])
def test_descontos_autonomos_conciliam_e_nao_duplicam_pago_no_vencimento(perfil, perfil_descontos):
    payload = entrada(perfil, perfil_descontos)
    original = copy.deepcopy(payload)
    original.pop('descontos')
    original.pop('custas_despesas')
    parcelas = calcular(original)
    desconto = copy.deepcopy(original)
    desconto['perfil'] = perfil_descontos or perfil
    desconto['parcelas'] = [{'numero': item['numero'], 'historico': item['descricao'], 'data_vencimento': item['data'], 'valor_bruto': item['valor']} for item in payload['descontos']['itens']]
    separado = calcular(desconto)
    r = calcular(payload)
    assert r.resumo.principal_apurado == D('800')
    assert r.resumo.valor_pago_na_data_parcelas == D('200')
    assert r.parcelas == parcelas.parcelas
    assert r.resumo.abatimentos == separado.resumo.total_atualizado
    assert r.resumo.total_atualizado == parcelas.resumo.total_atualizado - separado.resumo.total_atualizado + r.resumo.custas
    assert r.descontos.parcelas == separado.parcelas
    assert r.descontos.descontos is None
    assert r.descontos.custas_despesas == []
    assert len(r.memorias_abatimento) == 2
    assert sum(ab.valor_abatido for ab in r.memorias_abatimento) == r.resumo.abatimentos
    assert r.memorias_abatimento[1].saldo_anterior == r.memorias_abatimento[0].saldo_posterior
    assert r.premissas['entrada']['descontos']['itens'][0]['data'] == '2025-10-15'


def test_excedente_e_informado_e_nao_abate_custas():
    payload = entrada()
    payload['descontos']['itens'][0]['valor'] = '5000'
    r = calcular(payload)
    assert r.resumo.total_atualizado == r.resumo.custas
    assert r.resumo.abatimentos == sum(p.total_parcela for p in r.parcelas)
    assert sum(ab.saldo_remanescente for ab in r.memorias_abatimento) == r.descontos.resumo.total_atualizado - r.resumo.abatimentos
    assert D(r.premissas['descontos']['excedente']) > 0
    assert any('excedente não foi compensado' in alerta for alerta in r.alertas)


@pytest.mark.parametrize('alteracao,erro', [
    ({'data': '2026-09-01'}, 'posterior à data-base'),
    ({'data': '2009-06-30'}, '01/07/2009'),
    ({'valor': '0'}, 'greater than 0'),
    ({'valor': '-10'}, 'greater than 0'),
    ({'numero': 2}, 'números distintos'),
])
def test_validacao_descontos(alteracao, erro):
    payload = entrada()
    payload['descontos']['itens'][0].update(alteracao)
    with pytest.raises(ValidationError, match=erro):
        CalculoSimplificado.model_validate(payload)


def test_marco_de_juros_proprio_e_nunca_anterior_ao_desconto():
    payload = entrada(perfil_descontos='selic_ipcae_poupanca_v1')
    payload['descontos'].update(criterio_inicio_juros='citacao', data_inicial_juros='2026-01-01')
    r = calcular(payload)
    assert str(r.descontos.parcelas[0].componentes['poupanca_pos'].data_inicial) == '2026-01-01'
    assert str(r.descontos.parcelas[1].componentes['poupanca_pos'].data_inicial) == '2026-02-10'
    assert r.descontos.parcelas[0].componentes['ipcae_pos'].data_inicial.isoformat() == '2025-10-15'
    payload['descontos']['data_inicial_juros'] = None
    with pytest.raises(ValidationError, match='Descontos: informe a data'):
        CalculoSimplificado.model_validate(payload)
    payload['descontos']['perfil'] = 'selic_cjf_v1'
    assert calcular(payload).descontos.dados_gerais.criterio_inicio_juros == 'vencimento'


def test_sem_descontos_preserva_total_e_memoria():
    payload = entrada()
    payload['descontos']['itens'] = []
    r = calcular(payload)
    payload.pop('descontos')
    assert r == calcular(payload)
    assert r.descontos is None
    assert r.resumo.abatimentos == 0
    assert r.memorias_abatimento == []


@pytest.mark.parametrize('perfil_descontos', ['selic_cjf_v1', 'selic_ipcae_2aa_v1', 'ipcae_1am_simples_v1', 'selic_ipcae_poupanca_v1'])
def test_api_e_exportacoes_incluem_descontos_indices_memoria_e_totais(perfil_descontos):
    payload = entrada(perfil_descontos=perfil_descontos)
    client = TestClient(app)
    response = client.post('/api/v1/calculo', json=payload)
    assert response.status_code == 200, response.text
    r = response.json()
    assert r['descontos']['premissas']['perfil'] == perfil_descontos
    response = client.post('/api/v1/calculo/exportar/pdf', json=payload)
    assert response.status_code == 200, response.text
    reader = PdfReader(io.BytesIO(response.content))
    texto = ' '.join(' '.join(page.extract_text() for page in reader.pages).split())
    assert texto.count('Data do desconto') == 1
    assert 'Pagamento posterior' in texto
    assert 'Crédito posterior' in texto
    assert 'descontos abatidos' in texto
    for p in r['descontos']['parcelas']:
        assert f"{D(p['total_parcela']):.2f}".replace('.', ',') in texto
    custo_paginas = [page.extract_text() for page in reader.pages if 'Fator IPCA-E' in page.extract_text()]
    assert len(custo_paginas) == 1
    assert 'Descontos atualizados' not in custo_paginas[0]
    response = client.post('/api/v1/calculo/exportar/excel', json=payload)
    assert response.status_code == 200, response.text
    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    for nome in ['Descontos', 'Índices descontos', 'Memória descontos', 'Critérios descontos', 'Abatimentos']:
        assert nome in wb.sheetnames
    assert wb['Descontos']['C2'].value == 'Pagamento posterior'
    assert D(str(wb['Descontos']['F2'].value)) == D(r['memorias_abatimento'][0]['valor_abatido'])
    assert any(row[0] == '(-) Descontos atualizados abatidos' for row in wb['Resumo'].values)
    response = client.post('/api/v1/calculo/exportar/csv', json=payload)
    assert response.status_code == 200, response.text
    with zipfile.ZipFile(io.BytesIO(response.content)) as arquivo:
        for nome in ['descontos.csv', 'indices_descontos.csv', 'memoria_descontos.csv', 'premissas_descontos.json']:
            assert nome in arquivo.namelist()
        assert 'Pagamento posterior' in arquivo.read('descontos.csv').decode('utf-8-sig')
        assert json.loads(arquivo.read('premissas_descontos.json'))['perfil'] == perfil_descontos


def test_excel_nao_trunca_entrada_extensa_e_nao_interpreta_descricao_como_formula(tmp_path):
    payload = entrada()
    payload['descontos']['itens'][0]['descricao'] = '=texto literal'
    r = calcular(payload)
    r.premissas['entrada']['descontos']['itens'] = [{'numero': i + 1, 'descricao': 'x' * 240, 'data': '2025-10-15', 'valor': '1'} for i in range(2000)]
    destino = tmp_path / 'qa.xlsx'
    exportar_excel(r, str(destino))
    wb = openpyxl.load_workbook(destino)
    assert wb['Descontos']['C2'].data_type == 's'
    linhas = list(wb['Premissas'].values)
    assert sum(str(linha[0]).startswith('Entrada desconto ') for linha in linhas) == 2000
    assert all(len(str(linha[1])) < 32767 for linha in linhas)


def test_selecao_preserva_lancamento_mas_nao_aplica_encargos_ou_abatimento():
    payload = entrada('selic_cjf_v1')
    payload['descontos']['itens'][0].update(aplicar=False, data='2099-12-31')
    r = calcular(payload)
    assert [p.numero for p in r.descontos.parcelas] == [2]
    assert all(m.parcela == 2 for m in r.descontos.memoria_mensal)
    assert len(r.memorias_abatimento) == 1
    assert r.premissas['entrada']['descontos']['itens'][0]['data'] == '2099-12-31'
    assert r.premissas['selecao_descontos'][0]['aplicar'] is False
    assert r.premissas['descontos']['quantidade'] == 1


def test_todos_desmarcados_preservam_saldo_e_nao_exigem_cobertura_ou_juros():
    payload = entrada('selic_cjf_v1', 'ipca_taxa_legal_v1')
    payload['descontos'].update(criterio_inicio_juros='citacao')
    payload['descontos']['itens'] = [{'numero': 1, 'aplicar': False}]
    r = calcular(payload)
    payload.pop('descontos')
    sem = calcular(payload)
    assert r.resumo == sem.resumo
    assert r.parcelas == sem.parcelas
    assert r.memoria_mensal == sem.memoria_mensal
    assert r.descontos is None and r.resumo.abatimentos == 0
    assert r.premissas['selecao_descontos'][0]['valor'] is None


@pytest.mark.parametrize('perfil', ['selic_cjf_v1', 'selic_ipcae_poupanca_v1', 'selic_ipcae_2aa_v1', 'ipca_taxa_legal_v1'])
@pytest.mark.parametrize('nominal', ['45.82', '5845.82'])
def test_parte_nominal_nao_recebe_encargos_e_concilia_todos_os_componentes(perfil, nominal):
    payload = entrada('selic_cjf_v1', perfil)
    payload['parcelas'][0]['valor_bruto'] = '20000'
    payload['descontos']['itens'] = [{'numero': 1, 'data': '2025-12-31', 'valor': '5845.82', 'encargos_sem_atualizacao': nominal}]
    r = calcular(payload)
    base = D('5845.82') - D(nominal)
    simples = copy.deepcopy(payload)
    simples['perfil'] = perfil
    simples['parcelas'] = [{'numero': 1, 'data_vencimento': '2025-12-31', 'valor_bruto': '5845.82', 'valor_pago_na_data': nominal}]
    simples.pop('descontos')
    simples.pop('custas_despesas')
    separado = calcular(simples)
    assert r.descontos.resumo.total_atualizado == separado.resumo.total_atualizado + D(nominal)
    assert r.descontos.resumo.principal_original == D('5845.82')
    assert r.descontos.resumo.principal_apurado == D('5845.82')
    assert r.descontos.resumo.valor_pago_na_data_parcelas == 0
    assert r.descontos.parcelas[0].valor_pago_na_data == 0
    assert r.descontos.parcelas[0].valor_apurado == D('5845.82')
    assert sum(r.descontos.resumo.totais_componentes.values()) + D('5845.82') == r.descontos.resumo.total_atualizado
    assert [m.valor_base for m in r.descontos.memoria_mensal] == [m.valor_base for m in separado.memoria_mensal]
    assert r.descontos.parcelas[0].componentes == separado.parcelas[0].componentes
    assert r.resumo.abatimentos == r.descontos.resumo.total_atualizado


@pytest.mark.parametrize('item', [
    {'numero': 1, 'valor': '1'}, {'numero': 1, 'data': '2025-12-31'},
    {'numero': 1, 'data': '2025-12-31', 'valor': '10', 'encargos_sem_atualizacao': '10.01'},
    {'numero': 1, 'data': '2025-12-31', 'valor': '10', 'encargos_sem_atualizacao': '-1'},
])
def test_dados_ativos_e_composicao_invalidos_sao_rejeitados(item):
    payload = entrada()
    payload['descontos']['itens'] = [item]
    with pytest.raises(ValidationError):
        calcular(payload)


def test_caso_0072089_recalculado_desde_original_e_pagamento_sem_capitalizacao():
    payload = entrada('selic_cjf_v1')
    payload['dados_gerais']['data_base'] = '2026-09-01'
    payload['parcelas'] = [{'numero': 1, 'data_vencimento': '2024-11-14', 'valor_bruto': '5800'}]
    payload['custas_despesas'] = []
    payload['descontos']['itens'] = [{'numero': 1, 'aplicar': True, 'data': '2025-12-31', 'valor': '5845.82', 'encargos_sem_atualizacao': '45.82'}]
    r = calcular(payload)
    assert r.parcelas[0].total_parcela == D('7200.70')
    assert r.descontos.parcelas[0].componentes['selic'].base_calculo == D('5800')
    assert r.descontos.parcelas[0].componentes['selic'].taxa_acumulada_percentual == D('10.18')
    assert r.resumo.abatimentos == D('6436.26')
    assert r.resumo.total_atualizado == D('764.44')
    from liquidacao_custom.core.relatorio_pdf import exportar_pdf
    reader = PdfReader(io.BytesIO(exportar_pdf(r)))
    posicoes = []
    def registrar(texto, cm, tm, fonte, tamanho):
        if 'R$ 764,44' in texto and tamanho == 14:
            posicoes.append(cm[4] + tm[4])
    reader.pages[0].extract_text(visitor_text=registrar)
    assert posicoes and all(x > 650 for x in posicoes), 'Total deve ocupar a extremidade direita do resumo.'
    payload['descontos']['itens'][0]['aplicar'] = False
    assert calcular(payload).resumo.total_atualizado == D('7200.70')


def test_exportacoes_conservam_selecao_composicao_e_nao_duplicam_pagamento():
    payload = entrada('selic_cjf_v1')
    payload['descontos']['itens'][0].update(encargos_sem_atualizacao='45.82')
    payload['descontos']['itens'][1].update(aplicar=False, descricao='=Crédito não selecionado')
    client = TestClient(app)
    response = client.post('/api/v1/calculo/exportar/pdf', json=payload)
    assert response.status_code == 200, response.text
    reader = PdfReader(io.BytesIO(response.content))
    texto = ' '.join(' '.join(page.extract_text() for page in reader.pages).split())
    assert texto.count('Data do desconto') == 1
    assert 'Pagamentos não abatidos' in texto
    assert '104,18' in texto and '45,82' in texto
    assert 'sem nova incidência' in texto
    response = client.post('/api/v1/calculo/exportar/excel', json=payload)
    assert response.status_code == 200, response.text
    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    assert wb['Pagamentos registrados']['E2'].value == '104.18'
    assert wb['Pagamentos registrados']['F2'].value == '45.82'
    assert wb['Pagamentos registrados']['G3'].value == 'NÃO'
    assert wb['Pagamentos registrados']['B3'].data_type == 's'
    assert wb['Descontos'].max_row == 2
    response = client.post('/api/v1/calculo/exportar/csv', json=payload)
    assert response.status_code == 200, response.text
    with zipfile.ZipFile(io.BytesIO(response.content)) as arquivo:
        assert 'NÃO' in arquivo.read('pagamentos_registrados.csv').decode('utf-8-sig')
        assert 'Crédito não selecionado' not in arquivo.read('descontos.csv').decode('utf-8-sig')


@pytest.mark.parametrize('quantidade,perfil', [(20, 'selic_cjf_v1'), (36, 'selic_ipcae_poupanca_v1')])
def test_pdf_muitas_parcelas_com_selecao_composicao_e_textos_longos(quantidade, perfil):
    from liquidacao_custom.core.relatorio_pdf import exportar_pdf
    payload = entrada(perfil, 'selic_cjf_v1')
    payload['parcelas'] = [{**payload['parcelas'][0], 'numero': i + 1, 'historico': f'Parcela {i + 1}', 'data_vencimento': '2021-01-01'} for i in range(quantidade)]
    payload['dados_gerais']['observacoes'] = 'Observação extensa para conferir a preservação dos critérios e das margens. ' * 30
    payload['descontos']['itens'] = [{'numero': i + 1, 'aplicar': i % 3 != 0, 'descricao': (f'Pagamento {i + 1} com descrição longa. ' * 8)[:240], 'data': '2025-12-31', 'valor': '150', 'encargos_sem_atualizacao': '5'} for i in range(quantidade)]
    r = calcular(payload)
    reader = PdfReader(io.BytesIO(exportar_pdf(r)))
    texto = ' '.join(page.extract_text() for page in reader.pages)
    assert texto.count('Data do\ndesconto') == 1
    assert 'Pagamentos não abatidos' in texto and 'sem nova incidência' in texto
    assert len(r.descontos.parcelas) == sum(item['aplicar'] for item in payload['descontos']['itens'])
    assert r.resumo.total_atualizado == sum(p.total_parcela for p in r.parcelas) - r.resumo.abatimentos + r.resumo.custas
    for i, pagina in enumerate(reader.pages):
        assert round(float(pagina.mediabox.width)) == 842
        assert round(float(pagina.mediabox.height)) == 595
        assert f'Página {i + 1}' in pagina.extract_text()
    assert reader.metadata.title == 'Demonstrativo de cálculo'
    assert reader.metadata.author == 'Barreto Fontes Sociedade de Advogados'
