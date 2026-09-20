import copy
import io
import zipfile
from datetime import date, timedelta
from decimal import Decimal as D

import openpyxl
import pytest
from fastapi.testclient import TestClient
from pydantic import ValidationError
from pypdf import PdfReader

from liquidacao_custom.api.main import app
from liquidacao_custom.core import motor_simplificado as motor
from liquidacao_custom.core.criterios_simplificados import CalculoSimplificado
from liquidacao_custom.core.importacao_simplificada import COLUNAS, COLUNAS_ANTIGAS, importar_parcelas


def entrada(perfil="selic_ipcae_poupanca_v1", **parcela):
    return {"perfil": perfil, "dados_gerais": {"data_base": "2026-08-31", "criterio_inicio_juros": "vencimento"},
            "parcelas": [{"numero": 1, "historico": "Parcela com multa", "data_vencimento": "2025-10-01", "valor_bruto": "1000", "valor_pago_na_data": "200", "multa_percentual": "2", **parcela}]}


def calcular(payload):
    return motor.executar_calculo(CalculoSimplificado.model_validate(payload))


@pytest.mark.parametrize("perfil", ["selic_ipcae_poupanca_v1", "selic_ipcae_2aa_v1", "selic_cjf_v1", "ipcae_1am_simples_v1"])
def test_multa_recebe_mesmos_encargos_sem_duplicacao(perfil):
    payload = entrada(perfil, data_vencimento="2019-06-05")
    sem_multa = copy.deepcopy(payload)
    sem_multa["parcelas"][0]["multa_percentual"] = "0"
    principal = calcular(sem_multa)
    nominal = copy.deepcopy(sem_multa)
    nominal["parcelas"][0].update(valor_bruto="16", valor_pago_na_data="0")
    propria = calcular(nominal)
    r = calcular(payload)
    p = r.parcelas[0]
    m = p.multa_detalhes
    assert m.base_calculo == 800 and m.valor_original == 16
    assert m.juros_mora == propria.resumo.juros_mora
    assert m.correcao_monetaria == propria.resumo.correcao_monetaria
    assert m.componentes == propria.parcelas[0].componentes
    assert m.total_atualizado == propria.resumo.total_atualizado
    assert r.resumo.total_atualizado == principal.resumo.total_atualizado + m.total_atualizado
    assert sum((c.valor for c in p.componentes.values()), p.valor_apurado) == p.total_parcela
    assert r.resumo.principal_apurado + r.resumo.correcao_monetaria + r.resumo.juros_mora + r.resumo.multas == r.resumo.total_atualizado
    assert all(x.indice_aplicado.startswith("Multa - ") for x in m.memoria)
    assert len(r.memoria_mensal) == len(principal.memoria_mensal) + len(m.memoria)
    if perfil == "selic_cjf_v1":
        assert m.juros_mora == 0 and set(m.componentes) == {"selic"}


def test_juros_da_multa_numericos_e_inicio_informado(monkeypatch):
    dados = {"selic": {}, "ipcae_taxas": {"2025-10": "10", "2025-11": "10"},
             "poupanca_total": {}, "fontes": {}, "obtido_em": "2026-09-06"}
    dia = date(2025, 10, 1)
    while dia <= date(2025, 12, 1):
        dados["poupanca_total"][dia.isoformat()] = "0.5"
        dia += timedelta(days=1)
    monkeypatch.setattr(motor, "carregar_base", lambda: dados)
    payload = entrada(valor_pago_na_data="0")
    payload["dados_gerais"].update(data_base="2025-11-30", criterio_inicio_juros="data_fixa", data_inicial_juros="2025-11-01")
    r = calcular(payload)
    m = r.parcelas[0].multa_detalhes
    assert m.valor_original == D("20.00")
    assert m.correcao_monetaria == D("4.20")
    assert m.juros_mora == D("0.12")  # 24,20 × 0,5%, somente novembro.
    assert m.total_atualizado == D("24.32")
    assert m.componentes["poupanca_pos"].data_inicial == date(2025, 11, 1)
    assert m.componentes["poupanca_pos"].base_calculo == D("24.2")


def test_multas_individuais_quitacao_e_percentual_zero():
    payload = entrada()
    payload["parcelas"] += [{**payload["parcelas"][0], "numero": 2, "multa_percentual": "5"},
                             {**payload["parcelas"][0], "numero": 3, "multa_percentual": "0"},
                             {**payload["parcelas"][0], "numero": 4, "valor_pago_na_data": "1000"}]
    r = calcular(payload)
    assert [p.multa_detalhes.valor_original if p.multa_detalhes else None for p in r.parcelas] == [D(16), D(40), None, D(0)]
    assert r.parcelas[3].total_parcela == 0
    assert r.parcelas[2].componentes["multa_parcela"].data_inicial is None
    assert r.resumo.multas == sum(p.multa for p in r.parcelas)


@pytest.mark.parametrize("percentual,nominal", [("2", "2.01"), ("0.0001", "0.00"), ("100", "100.25")])
def test_arredondamento_e_limites_sem_encargo_no_mes_atual(percentual, nominal):
    r = calcular(entrada("selic_cjf_v1", data_vencimento="2026-08-01",
                         valor_bruto="100.25", valor_pago_na_data="0", multa_percentual=percentual))
    m = r.parcelas[0].multa_detalhes
    assert m.valor_original == D(nominal)
    assert m.total_atualizado == D(nominal)
    assert r.resumo.total_atualizado == D("100.25") + D(nominal)
    assert any("Total da parcela = principal + SELIC + multa e encargos." in nota
               for nota in r.premissas["metodologia"])


def test_multas_antes_de_descontos_e_fora_das_custas():
    payload = entrada()
    payload["descontos"] = {"perfil": "selic_cjf_v1", "itens": [{"numero": 1, "data": "2026-02-01", "valor": "50"}]}
    payload["custas_despesas"] = [{"numero": 1, "nome": "Taxa", "data": "2026-02-01", "valor": "30"}]
    r = calcular(payload)
    assert r.parcelas[0].multa_detalhes.base_calculo == 800
    assert r.resumo.total_atualizado == sum(p.total_parcela for p in r.parcelas) - r.resumo.abatimentos + r.resumo.custas
    assert r.descontos.resumo.multas == 0


@pytest.mark.parametrize("valor", ["-1", "100.0001", "NaN", "Infinity"])
def test_percentual_invalido_rejeitado(valor):
    with pytest.raises(ValidationError):
        CalculoSimplificado.model_validate(entrada(multa_percentual=valor))


def test_sem_multa_preserva_resultado():
    payload = entrada(multa_percentual="0")
    r = calcular(payload)
    del payload["parcelas"][0]["multa_percentual"]
    assert calcular(payload).model_dump() == r.model_dump()
    assert r.parcelas[0].multa_detalhes is None
    assert "multa_parcela" not in r.resumo.totais_componentes


@pytest.mark.parametrize("legado", [False, True])
def test_importacao_preserva_multa_e_aceita_modelo_antigo(tmp_path, legado):
    wb = openpyxl.Workbook()
    wb.active.append(COLUNAS_ANTIGAS if legado else COLUNAS)
    wb.active.append([1, "Parcela", "2025-10-01", 1000, 200, None] + ([] if legado else [2]))
    caminho = tmp_path / "entrada.xlsx"
    wb.save(caminho)
    wb.close()
    parcelas, erros = importar_parcelas(caminho)
    assert not erros and parcelas[0].multa_percentual == (0 if legado else 2)


def test_api_e_exportacoes_discriminam_multa():
    client = TestClient(app)
    payload = entrada()
    response = client.post("/api/v1/calculo", json=payload)
    assert response.status_code == 200
    assert D(response.json()["parcelas"][0]["multa_detalhes"]["juros_mora"]) > 0
    pdf = client.post("/api/v1/calculo/exportar/pdf", json=payload)
    assert pdf.status_code == 200
    reader = PdfReader(io.BytesIO(pdf.content))
    content = " ".join(" ".join(page.extract_text().split()) for page in reader.pages)
    assert "Multas por parcela" in content and "Multa nominal" in content
    assert content.count("Parcela / descrição") == 2  # Uma vez em cada tabela distinta.
    assert "sem nova capitalização" in content
    excel = client.post("/api/v1/calculo/exportar/excel", json=payload)
    assert excel.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(excel.content))
    assert "Multas parcelas" in wb.sheetnames
    assert wb["Multas parcelas"]["E2"].value == 2
    wb.close()
    csv = client.post("/api/v1/calculo/exportar/csv", json=payload)
    assert csv.status_code == 200
    with zipfile.ZipFile(io.BytesIO(csv.content)) as archive:
        assert "multas_parcelas.csv" in archive.namelist()
        assert "Multa - Poupança" in archive.read("indices_por_parcela.csv").decode("utf-8-sig")


def test_resumo_pdf_com_multa_e_custas_mantem_total_na_extremidade_direita(monkeypatch):
    from liquidacao_custom.core import relatorio_pdf
    tabela_original = relatorio_pdf.Table
    tabelas = []

    def registrar(dados, *args, **kwargs):
        tabelas.append(dados)
        return tabela_original(dados, *args, **kwargs)

    monkeypatch.setattr(relatorio_pdf, "Table", registrar)
    payload = entrada()
    payload["custas_despesas"] = [{"numero": 1, "nome": "Taxa", "data": "2025-10-01", "valor": "100"}]
    relatorio_pdf.exportar_pdf(calcular(payload))
    resumo = next(dados for dados in tabelas
                  if len(dados) == 6 and len(dados[0]) == 4)
    assert resumo[-2][:3] == ["", "", ""]
    assert resumo[-2][-1].getPlainText() == "TOTAL ATUALIZADO"
    assert resumo[-1][-1].getPlainText() == "R$ 1.012,35"
