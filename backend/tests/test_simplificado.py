import io
import json
import zipfile
from datetime import date, timedelta
from decimal import Decimal as D

import openpyxl
import pytest
from fastapi.testclient import TestClient
from pydantic import ValidationError

from liquidacao_custom.api.main import app
from liquidacao_custom.core.criterios_simplificados import CalculoSimplificado
from liquidacao_custom.core import motor_simplificado as motor
from liquidacao_custom.core.importacao_simplificada import gerar_template, importar_parcelas


def entrada(inicio="2025-09-01", fim="2025-09-30", **parcela):
    return {"dados_gerais": {"data_base": fim}, "parcelas": [{
        "numero": 1, "data_vencimento": inicio, "valor_bruto": "1000.00", **parcela,
    }]}


@pytest.fixture
def indices(monkeypatch):
    data = {
        "selic": {"2025-09": "3", "2025-01": "1", "2025-02": "2", "2024-02": "2.9", "2021-12": "3.1"},
        "ipcae_taxas": {"2025-09": "3", "2025-10": "0", "2025-11": "0", "2025-12": "0"},
        "poupanca_total": {}, "fontes": {"ipcae": "https://example.test/ipca15"}, "obtido_em": "2026-09-06",
    }
    atual = date(2025, 9, 1)
    while atual <= date(2026, 1, 1):
        data["poupanca_total"][atual.isoformat()] = "0.5"
        atual += timedelta(days=1)
    monkeypatch.setattr(motor, "carregar_base", lambda: data)
    return data


def calcular(**kwargs):
    return motor.executar_calculo(CalculoSimplificado.model_validate(entrada(**kwargs)))


def test_transicao_sem_lacuna_ou_sobreposicao(indices):
    r = calcular()
    assert r.resumo.total_atualizado == D("1033.70")
    assert r.resumo.juros_mora == D("3.61")
    assert [m.indice_aplicado for m in r.memoria_mensal] == ["SELIC única", "IPCA-E IBGE (IPCA-15)", "Juros da poupança"]
    assert "2025-09-01 a 2025-09-09" in r.memoria_mensal[0].observacao
    assert "2025-09-10 a 2025-09-30" in r.memoria_mensal[1].observacao
    assert r.resumo.principal_apurado + r.resumo.correcao_monetaria + r.resumo.juros_mora == r.resumo.total_atualizado


def test_selic_soma_taxas_sem_capitalizar(indices):
    r = calcular(inicio="2025-01-01", fim="2025-02-28")
    assert r.resumo.total_atualizado == D("1030.00")
    assert r.resumo.juros_mora == 0
    c = r.parcelas[0].componentes
    assert c["selic"].taxa_acumulada_percentual == D("3")
    assert c["selic"].valor == D("30.00")
    assert c["selic"].base_calculo == D("1000")
    assert c["ipcae_pre"].data_inicial is None
    assert c["ipcae_pos"].data_inicial is None


@pytest.mark.parametrize("inicio,fim,total", [
    ("2021-12-09", "2021-12-31", "1023.00"),
    ("2024-02-29", "2024-02-29", "1001.00"),
    ("2025-09-09", "2025-09-09", "1001.00"),
])
def test_dias_limite_e_ano_bissexto(indices, inicio, fim, total):
    assert calcular(inicio=inicio, fim=fim).resumo.total_atualizado == D(total)


def test_primeiro_dia_novo_criterio_nao_aplica_selic(indices):
    r = calcular(inicio="2025-09-10", fim="2025-09-10")
    assert all("SELIC" not in m.indice_aplicado for m in r.memoria_mensal)
    assert len(r.memoria_mensal) == 2
    assert r.resumo.juros_mora == D("0.17")


def test_ipca15_meses_completos_e_poupanca_simples(indices):
    indices["ipcae_taxas"].update({"2025-10": "10", "2025-11": "10"})
    r = calcular(inicio="2025-10-01", fim="2025-11-30")
    assert r.resumo.correcao_monetaria == D("210.00")
    assert r.resumo.juros_mora == D("12.10")
    assert r.resumo.total_atualizado == D("1222.10")


def test_custas_despesas_usam_apenas_ipcae_e_somam_somente_ao_total_final(indices):
    sem_custas = calcular(inicio="2025-09-10", fim="2025-09-30")
    dados = entrada(inicio="2025-09-10", fim="2025-09-30")
    dados["custas_despesas"] = [{
        "numero": 1,
        "nome": "Honorários periciais",
        "data": "2025-09-01",
        "valor": "100.00",
    }]

    resultado = motor.executar_calculo(CalculoSimplificado.model_validate(dados))
    custa = resultado.custas_despesas[0]

    assert custa.nome == "Honorários periciais"
    assert custa.fator_ipcae == D("1.03")
    assert custa.correcao_monetaria == D("3.00")
    assert custa.valor_atualizado == D("103.00")
    assert all(memoria.juros_periodo == 0 for memoria in custa.memoria)
    assert resultado.resumo.principal_apurado == sem_custas.resumo.principal_apurado
    assert resultado.resumo.correcao_monetaria == sem_custas.resumo.correcao_monetaria
    assert resultado.resumo.juros_mora == sem_custas.resumo.juros_mora
    assert resultado.resumo.custas == D("103.00")
    assert resultado.resumo.total_atualizado == sem_custas.resumo.total_atualizado + D("103.00")
    assert resultado.premissas["custas_despesas"]["criterio"].startswith("IPCA-E")


def test_custas_despesas_aparecem_no_pdf_excel_e_csv(indices):
    dados = entrada(inicio="2025-09-10", fim="2025-09-30")
    dados["custas_despesas"] = [{
        "numero": 1,
        "nome": "Taxa judiciária",
        "data": "2025-09-01",
        "valor": "150.00",
    }]
    cliente = TestClient(app)

    pdf = cliente.post("/api/v1/calculo/exportar/pdf", json=dados)
    assert pdf.status_code == 200, pdf.text
    from pypdf import PdfReader
    paginas_pdf = PdfReader(io.BytesIO(pdf.content)).pages
    texto_pdf = "\n".join(page.extract_text() for page in paginas_pdf)
    primeira_pagina = paginas_pdf[0].extract_text()
    pagina_custas = paginas_pdf[1].extract_text()
    assert "Custas e despesas processuais" in texto_pdf
    assert "Taxa judiciária" in texto_pdf
    assert "Fator IPCA-E" in texto_pdf
    assert "Taxa judiciária" not in primeira_pagina
    assert "Custas e despesas processuais" in pagina_custas
    assert "Taxa judiciária" in pagina_custas
    assert "Valores por parcela" not in pagina_custas

    excel = cliente.post("/api/v1/calculo/exportar/excel", json=dados)
    assert excel.status_code == 200, excel.text
    planilha = openpyxl.load_workbook(io.BytesIO(excel.content), data_only=True)
    assert "Custas e despesas" in planilha.sheetnames
    assert planilha["Custas e despesas"]["B2"].value == "Taxa judiciária"

    csv_zip = cliente.post("/api/v1/calculo/exportar/csv", json=dados)
    assert csv_zip.status_code == 200, csv_zip.text
    with zipfile.ZipFile(io.BytesIO(csv_zip.content)) as arquivo:
        assert "custas_despesas.csv" in arquivo.namelist()
        assert "memoria_custas_despesas.csv" in arquivo.namelist()
        assert "Taxa judiciária" in arquivo.read("custas_despesas.csv").decode("utf-8-sig")


def test_inicio_juros_posterior_e_pagamento_no_vencimento(indices):
    r = calcular(inicio="2025-10-01", fim="2025-11-30", data_inicial_juros="2025-11-01", valor_pago_na_data="200.00")
    assert r.resumo.principal_apurado == D("800")
    assert r.resumo.juros_mora == D("4.00")
    assert r.resumo.total_atualizado == D("804.00")


def test_poupanca_usa_taxa_total_publicada_e_nao_constante(indices):
    indices["poupanca_total"]["2025-09-10"] = "0"
    assert calcular(inicio="2025-09-10", fim="2025-09-10").resumo.juros_mora == 0
    indices["poupanca_total"]["2025-09-10"] = "0.75"
    assert calcular(inicio="2025-09-10", fim="2025-09-10").resumo.juros_mora == D("0.25")


def test_deflacao_preservada(indices):
    indices["ipcae_taxas"]["2025-10"] = "-2.912621359223300970873786408"
    r = calcular(inicio="2025-10-01", fim="2025-10-31", data_inicial_juros="2025-12-01")
    assert r.resumo.correcao_monetaria == D("-29.13")
    assert r.resumo.juros_mora == 0


@pytest.mark.parametrize("serie,chave", [("selic", "2025-09"), ("ipcae_taxas", "2025-09"), ("poupanca_total", "2025-09-10")])
def test_indice_ausente_nao_vira_zero(indices, serie, chave):
    del indices[serie][chave]
    with pytest.raises(ValueError):
        calcular()


@pytest.mark.parametrize("mudanca", [
    {"data_vencimento": "2009-06-30"}, {"data_vencimento": "2026-01-01"},
    {"valor_bruto": "-1"}, {"valor_pago_na_data": "1001"},
    {"correcao_monetaria": "ipca"},
])
def test_entrada_invalida(mudanca):
    with pytest.raises(ValidationError):
        CalculoSimplificado.model_validate(entrada(**mudanca))


@pytest.mark.parametrize("endpoint", ["calculo", "calculo/exportar/excel", "calculo/exportar/csv", "calculo/exportar/pdf"])
def test_api_rejeita_criterios_legados_em_todas_saidas(endpoint):
    payload = entrada()
    payload["config_selic_taxa_unica"] = {"aplicar": False}
    assert TestClient(app).post("/api/v1/" + endpoint, json=payload).status_code == 422


def test_exportacoes_preservam_entrada_e_premissas():
    client = TestClient(app)
    payload = entrada()
    esperado = client.post("/api/v1/calculo", json=payload).json()
    excel = client.post("/api/v1/calculo/exportar/excel", json=payload)
    assert excel.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(excel.content))
    premissas = dict(wb["Premissas"].iter_rows(min_row=2, values_only=True))
    assert json.loads(premissas["entrada"])["dados_gerais"] == esperado["premissas"]["entrada"]["dados_gerais"]
    assert json.loads(premissas["Entrada parcela 1"]) == esperado["premissas"]["entrada"]["parcelas"][0]
    assert set(wb.sheetnames) == {"Dados Gerais", "Parcelas", "Memória Mensal", "Premissas", "Resumo", "Índices por parcela"}
    assert json.loads(premissas["sha256_base"]) == esperado["premissas"]["sha256_base"]
    resumo = dict(wb["Resumo"].iter_rows(min_row=2, values_only=True))
    assert D(str(resumo["TOTAL ATUALIZADO DO DÉBITO"])) == D(esperado["resumo"]["total_atualizado"])
    wb.close()
    csv = client.post("/api/v1/calculo/exportar/csv", json=payload)
    assert csv.status_code == 200
    with zipfile.ZipFile(io.BytesIO(csv.content)) as z:
        assert json.loads(z.read("premissas.json")) == esperado["premissas"]
        import csv as csv_lib
        linhas = list(csv_lib.reader(io.StringIO(z.read("indices_por_parcela.csv").decode("utf-8-sig"))))
        assert len(linhas) == 6  # Cabeçalho e os cinco componentes da parcela.
        assert next(linha for linha in linhas if linha[1] == "SELIC única")[7] == f"{D(esperado['parcelas'][0]['componentes']['selic']['taxa_acumulada_percentual']):.4f}"


def test_modelo_simplificado_importa_e_rejeita_planilha_antiga(tmp_path):
    caminho = tmp_path / "modelo.xlsx"
    gerar_template(caminho)
    parcelas, erros = importar_parcelas(caminho)
    assert len(parcelas) == 1 and not erros
    wb = openpyxl.load_workbook(caminho)
    wb.active.cell(1, 7, "correcao_monetaria")
    wb.save(caminho)
    wb.close()
    assert importar_parcelas(caminho)[1]


def test_base_oficial_e_cobertura():
    data = motor.carregar_base()
    assert data["ipcae_taxas"]["2025-09"] == "0.48"
    assert data["selic"]["2025-09"] == "1.22"
    assert motor.criterios_publicos()["data_base_maxima"] == "2026-08-31"
    r = calcular(inicio="2021-12-09", fim="2026-08-31")
    assert not r.alertas and r.resumo.total_atualizado > 1000
    with pytest.raises(ValueError, match="31/08/2026"):
        calcular(inicio="2025-09-10", fim="2026-09-06")


def test_citacao_anterior_nao_antecipa_juros_ao_vencimento(indices):
    r = calcular(inicio="2025-10-01", fim="2025-10-31", data_inicial_juros="2024-01-01")
    assert r.resumo.juros_mora == D("5.00")


def test_nao_muta_entrada_e_reexecuta_com_mesmo_resultado(indices):
    c = CalculoSimplificado.model_validate(entrada())
    antes = c.model_dump_json()
    r1 = motor.executar_calculo(c)
    assert c.model_dump_json() == antes
    assert motor.executar_calculo(c) == r1


@pytest.mark.parametrize("endpoint", ["calculo", "calculo/exportar/excel", "calculo/exportar/csv", "calculo/exportar/pdf"])
@pytest.mark.parametrize("fim", ["2026-09-01", "2026-09-06", "2027-01-01"])
def test_api_bloqueia_apos_ultimo_mes_completo(endpoint, fim):
    response = TestClient(app).post("/api/v1/" + endpoint, json=entrada(fim=fim))
    assert response.status_code == 400
    assert "data-base máxima permitida é 31/08/2026" in response.json()["detail"]


@pytest.mark.parametrize("fim", ["2026-08-30", "2026-08-31"])
def test_api_aceita_datas_ate_limite(fim):
    response = TestClient(app).post("/api/v1/calculo", json=entrada(fim=fim))
    assert response.status_code == 200
    assert response.json()["dados_gerais"]["data_base"] == fim
    assert response.json()["premissas"]["data_base_maxima"] == "2026-08-31"


def test_limite_considera_cobertura_da_poupanca(indices):
    # A tabela tem fator para dezembro, mas novembro não tem os juros completos.
    del indices["poupanca_total"]["2025-11-15"]
    assert motor.criterios_publicos()["data_base_maxima"] == "2025-10-31"
    with pytest.raises(ValueError, match="31/10/2025"):
        calcular(fim="2025-11-01")


def test_limite_so_avanca_com_mes_inteiro_disponivel(indices):
    del indices["ipcae_taxas"]["2025-12"]
    assert motor.ultima_data_disponivel(indices) == date(2025, 11, 30)
    indices["ipcae_taxas"]["2025-12"] = "1"
    del indices["poupanca_total"]["2026-01-01"]
    assert motor.ultima_data_disponivel(indices) == date(2025, 11, 30)
    indices["poupanca_total"]["2026-01-01"] = "0.5"
    assert motor.ultima_data_disponivel(indices) == date(2025, 12, 31)


def test_bloqueio_nao_depende_do_saldo_ou_do_periodo_da_parcela():
    with pytest.raises(ValueError, match="31/08/2026"):
        calcular(inicio="2025-01-01", fim="2026-09-01", valor_pago_na_data="1000.00")


def test_tr_incluida_uma_unica_vez_na_taxa_total(indices):
    # Total composto publicado para TR 0,2% e adicional 0,5%: 0,701%.
    # Não pode ser 0,5% (sem TR), 0,7% (mera soma) nem 0,901% (TR duplicada).
    for ref in indices["poupanca_total"]:
        indices["poupanca_total"][ref] = "0.701"
    r = calcular(inicio="2025-10-01", fim="2025-10-31")
    assert r.resumo.juros_mora == D("7.01")
    assert r.resumo.correcao_monetaria == 0
    assert r.resumo.total_atualizado == D("1007.01")
    assert "remuneração total" in r.parcelas[0].memoria_juros[0].observacao
    assert "tema905_ipcae_ibge_poupanca_total_v4" == r.premissas["versao_metodologia"]


def test_poupanca_fim_de_mes_usa_referencia_seguinte(indices):
    indices["poupanca_total"]["2025-11-01"] = "0.93"
    r = calcular(inicio="2025-10-29", fim="2025-10-31")
    # Três dias / 31 de taxa mensal 0,93%, sem capitalização.
    assert r.resumo.juros_mora == D("0.90")
    assert "ref. 2025-11-01" in r.parcelas[0].memoria_juros[0].observacao
    del indices["poupanca_total"]["2025-11-01"]
    with pytest.raises(ValueError):
        calcular(inicio="2025-10-29", fim="2025-10-31")


def test_snapshot_total_corresponde_fonte_bcb():
    from pathlib import Path
    fonte = json.loads((motor.BASE.parent / "fontes" / "bcb_poupanca_total.json").read_text(encoding="utf-8-sig"))
    data = motor.carregar_base()
    from datetime import datetime
    for item in fonte:
        chave = datetime.strptime(item["data"], "%d/%m/%Y").date().isoformat()
        assert data["poupanca_total"][chave] == item["valor"]
        assert data["poupanca_periodos"][chave] == datetime.strptime(item["dataFim"], "%d/%m/%Y").date().isoformat()
    # Pontos conferidos com a tabela de remuneração total do BCB.
    assert data["poupanca_total"]["2025-09-01"] == "0.6751"
    assert data["poupanca_total"]["2026-08-01"] == "0.6701"
    assert "meta_selic_diaria" not in data


def test_snapshot_ipcae_corresponde_serie_ipca15_ibge():
    from datetime import datetime

    fonte = json.loads(
        (motor.BASE.parent / "fontes" / "bcb_ipca15_2009_2026.json").read_text(encoding="utf-8-sig")
    )
    dados = motor.carregar_base()

    assert len(fonte) == 206
    assert len(dados["ipcae_taxas"]) == len(fonte)
    for item in fonte:
        chave = datetime.strptime(item["data"], "%d/%m/%Y").strftime("%Y-%m")
        assert dados["ipcae_taxas"][chave] == item["valor"]

    assert dados["ipcae_taxas"]["2009-07"] == "0.22"
    assert dados["ipcae_taxas"]["2026-08"] == "-0.40"
    assert "ipcae_fatores" not in dados
    assert dados["fontes"]["ipcae_serie_bcb"].endswith("method=consultarGraficoPorId")


@pytest.fixture
def tema905(indices):
    indices["ipcae_taxas"].update({"2021-11": "10", "2021-12": "10"})
    atual = date(2021, 11, 1)
    while atual <= date(2021, 12, 8):
        indices["poupanca_total"][atual.isoformat()] = "0.701"
        atual += timedelta(days=1)
    return indices


def test_tema905_mes_completo_com_tr_sem_taxa_fixa(tema905):
    r = calcular(inicio="2021-11-01", fim="2021-11-30")
    assert r.resumo.correcao_monetaria == D("100.00")
    assert r.resumo.juros_mora == D("7.71")  # 1100 * 0,701%
    assert r.resumo.total_atualizado == D("1107.71")
    assert all("SELIC" not in m.indice_aplicado for m in r.memoria_mensal)


def test_transicao_tema905_selic_sem_sobrepor_e_sem_capitalizar_poupanca(tema905):
    r = calcular(inicio="2021-12-01", fim="2021-12-31")
    base_selic = D(1000) * D("1.1") ** (D(8) / 31)
    corrigido = base_selic * D("1.023")  # 23 dias de SELIC de 3,1% ao mês
    juros = corrigido * D("0.00701") * 8 / 31
    assert r.resumo.correcao_monetaria == motor.moeda(corrigido) - 1000
    assert r.resumo.juros_mora == motor.moeda(juros)
    assert r.parcelas[0].memoria_correcao[1].valor_base == base_selic
    assert "2021-12-01 a 2021-12-08" in r.parcelas[0].memoria_correcao[0].observacao
    assert "2021-12-09 a 2021-12-31" in r.parcelas[0].memoria_correcao[1].observacao
    assert len(r.parcelas[0].memoria_juros) == 1
    assert "2021-12-01 a 2021-12-08" in r.parcelas[0].memoria_juros[0].observacao
    c = r.parcelas[0].componentes
    assert c["ipcae_pre"].fator_acumulado == D("1.1") ** (D(8) / 31)
    assert c["ipcae_pre"].valor == motor.moeda(base_selic) - 1000
    assert c["selic"].taxa_acumulada_percentual == D("2.3")
    assert c["selic"].base_calculo == base_selic
    assert c["selic"].valor == motor.moeda(corrigido) - motor.moeda(base_selic)
    assert abs(c["poupanca_pre"].taxa_acumulada_percentual - D("0.701") * 8 / 31) < D("1e-24")
    assert c["poupanca_pre"].base_calculo == corrigido
    assert c["ipcae_pre"].data_final == date(2021, 12, 8)
    assert c["selic"].data_inicial == date(2021, 12, 9)
    assert c["ipcae_pos"].data_inicial is None


def test_tema905_juros_respeitam_data_inicial(tema905):
    r = calcular(inicio="2021-11-01", fim="2021-11-30", data_inicial_juros="2021-11-16")
    assert r.resumo.juros_mora == D("3.86")
    assert calcular(inicio="2021-11-01", fim="2021-11-30", data_inicial_juros="2021-12-09").resumo.juros_mora == 0


@pytest.mark.parametrize("serie,chave", [("ipcae_taxas", "2021-11"), ("poupanca_total", "2021-11-15")])
def test_tema905_nao_completa_indice_ausente(tema905, serie, chave):
    del tema905[serie][chave]
    with pytest.raises(ValueError, match="Índice ausente"):
        calcular(inicio="2021-11-01", fim="2021-11-30")


def test_ultimo_dia_tema905_e_primeiro_dia_selic(tema905):
    r = calcular(inicio="2021-12-08", fim="2021-12-08")
    assert len(r.parcelas[0].memoria_correcao) == 1
    assert "Tema 905" in r.parcelas[0].memoria_correcao[0].indice_aplicado
    assert r.resumo.juros_mora > 0
    assert calcular(inicio="2021-12-09", fim="2021-12-09").resumo.total_atualizado == D("1001.00")


def test_tema905_cobertura_historica_oficial_sem_lacunas():
    data = motor.carregar_base()
    atual = date(2009, 7, 1)
    while atual <= date(2021, 12, 8):
        assert motor.referencia_poupanca(atual) in data["poupanca_total"]
        assert atual.strftime("%Y-%m") in data["ipcae_taxas"]
        atual += timedelta(days=1)
    assert data["poupanca_total"]["2009-07-01"] == "0.6056"
    assert data["poupanca_total"]["2012-05-04"] == "0.5146"
    assert data["ipcae_taxas"]["2021-03"] == "0.93"
    assert motor.criterios_publicos()["inicio"] == "2009-07-01"


@pytest.mark.parametrize("inicio", ["2009-07-01", "2012-05-01", "2021-03-01"])
def test_calculo_atravessa_tres_periodos_com_dados_oficiais(inicio):
    r = calcular(inicio=inicio, fim="2026-08-31")
    assert len(r.premissas["criterios"]) == 3
    assert r.resumo.principal_apurado + r.resumo.correcao_monetaria + r.resumo.juros_mora == r.resumo.total_atualizado
    assert any("Tema 905" in m.indice_aplicado for m in r.parcelas[0].memoria_correcao)
    assert any("SELIC" in m.indice_aplicado for m in r.parcelas[0].memoria_correcao)
    assert all("2021-12-09" not in m.observacao for m in r.parcelas[0].memoria_juros)
    p = r.parcelas[0]
    c = p.componentes
    assert sum((x.valor for x in c.values()), p.valor_apurado) == p.total_parcela
    assert sum(c[k].valor for k in ("ipcae_pre", "selic", "ipcae_pos")) == p.correcao_monetaria
    assert c["poupanca_pre"].valor + c["poupanca_pos"].valor == p.juros_mora
    assert c["ipcae_pos"].data_inicial == date(2025, 9, 10)
    assert c["selic"].data_final == date(2025, 9, 9)
    assert r.resumo.totais_componentes == {k: v.valor for k, v in c.items()}


def test_componentes_pos_transicao_fator_taxa_e_centavos(indices):
    p = calcular().parcelas[0]
    c = p.componentes
    assert c["selic"].taxa_acumulada_percentual == D("0.9")
    assert c["selic"].valor == D("9")
    assert c["ipcae_pos"].base_calculo == D("1009")
    assert c["ipcae_pos"].fator_acumulado == D("1.03") ** D("0.7")
    assert c["ipcae_pos"].valor == D("21.09")
    assert abs(c["poupanca_pos"].taxa_acumulada_percentual - D("0.35")) < D("1e-24")
    assert c["poupanca_pos"].valor == D("3.61")
    assert sum((v.valor for v in c.values()), p.valor_apurado) == D("1033.70")


def test_componentes_distinguem_nao_incidente_de_taxa_zero_e_saldo_quitado(indices):
    indices["poupanca_total"]["2025-09-10"] = "0"
    c = calcular(inicio="2025-09-10", fim="2025-09-10", valor_pago_na_data="1000").parcelas[0].componentes
    assert c["selic"].data_inicial is None and c["selic"].taxa_acumulada_percentual is None
    assert c["poupanca_pos"].data_inicial == date(2025, 9, 10)
    assert c["poupanca_pos"].taxa_acumulada_percentual == 0
    assert c["ipcae_pos"].fator_acumulado > 1
    assert all(v.valor == 0 for v in c.values())


def test_componentes_preservam_deflacao_e_juros_posteriores(indices):
    indices["ipcae_taxas"]["2025-10"] = "-2.912621359223300970873786408"
    c = calcular(inicio="2025-10-01", fim="2025-10-31", data_inicial_juros="2025-12-01").parcelas[0].componentes
    assert c["ipcae_pos"].fator_acumulado == D(100) / 103
    assert c["ipcae_pos"].valor == D("-29.13")
    assert c["poupanca_pos"].data_inicial is None


@pytest.mark.parametrize("endpoint", ["calculo", "calculo/exportar/excel", "calculo/exportar/csv", "calculo/exportar/pdf"])
def test_api_aceita_tema905_em_todas_as_saidas(endpoint):
    response = TestClient(app).post("/api/v1/" + endpoint, json=entrada(inicio="2021-03-01"))
    assert response.status_code == 200


def test_limite_dezembro2021_exige_dois_regimes(tema905):
    dados = {**tema905, "selic": {"2021-12": "3.1"}, "ipcae_taxas": {
        "2021-11": "10", "2021-12": "10",
    }}
    assert motor.ultima_data_disponivel(dados) == date(2021, 12, 31)
    del dados["poupanca_total"]["2021-12-08"]
    assert motor.ultima_data_disponivel(dados) == date(2021, 11, 30)
