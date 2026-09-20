"""Planilha de entrada sem opções de índices ou acessórios."""
from datetime import datetime
import openpyxl
from .criterios_simplificados import ParcelaSimplificada

COLUNAS_ANTIGAS = ["numero", "historico", "data_vencimento", "valor_bruto", "valor_pago_na_data", "data_inicial_juros"]
COLUNAS = [*COLUNAS_ANTIGAS, "multa_percentual"]


def gerar_template(caminho):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Parcelas"
    ws.append(COLUNAS)
    ws.append([1, "Parcela de exemplo", "2025-09-01", 1000, 0, "2025-09-10", 0])
    ws.freeze_panes = "A2"
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 25
    wb.save(caminho)
    wb.close()


def importar_parcelas(caminho):
    wb = openpyxl.load_workbook(caminho, data_only=False, read_only=True)
    parcelas, erros, numeros = [], [], set()
    try:
        linhas = wb.active.iter_rows(values_only=True)
        cabecalho = next(linhas, ())
        if list(cabecalho) not in (COLUNAS, COLUNAS_ANTIGAS):
            return [], ["Use o modelo simplificado: " + ", ".join(COLUNAS) + ". Planilhas com critérios antigos não são aceitas."]
        for numero_linha, valores in enumerate(linhas, 2):
            if all(v is None for v in valores):
                continue
            if numero_linha > 2001:
                erros.append("Limite de 2.000 parcelas excedido.")
                break
            entrada = {k: (v.date() if isinstance(v, datetime) else v) for k, v in zip(cabecalho, valores) if v is not None}
            try:
                p = ParcelaSimplificada.model_validate(entrada)
                if p.numero in numeros:
                    raise ValueError("Número da parcela repetido.")
                numeros.add(p.numero)
                parcelas.append(p)
            except ValueError as exc:
                erros.append(f"Linha {numero_linha}: {exc}")
        if not parcelas and not erros:
            erros.append("A planilha não contém parcelas.")
        return parcelas, erros
    finally:
        wb.close()
