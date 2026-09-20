"""
Módulo de importação de parcelas a partir de modelo do Excel.
"""

from __future__ import annotations

import os
from datetime import datetime, date
from decimal import Decimal
import openpyxl

from liquidacao_custom.core.models import Parcela, IndiceCorrecao, TipoJuros


def gerar_template(caminho: str) -> None:
    """Gera o arquivo de modelo Excel para importação."""
    diretorio = os.path.dirname(caminho)
    if diretorio and not os.path.exists(diretorio):
        os.makedirs(diretorio)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Modelo de Importação"

    colunas = [
        "numero", "data_vencimento", "natureza", "historico",
        "valor_bruto", "valor_pago_na_data", "correcao_monetaria",
        "data_inicial_correcao", "juros_moratorios", "data_inicial_juros",
        "multa_moratoria", "observacao"
    ]

    for col_num, col_name in enumerate(colunas, 1):
        ws.cell(row=1, column=col_num, value=col_name)

    # Adicionar exemplo de linha
    exemplo = [
        1, "2024-01-10", "Aluguel", "Aluguel ref Jan/2024",
        1500.00, 0.00, "igpm", "2024-01-10",
        "percentual_mensal_simples", "2024-01-10", 10.00, "Primeira parcela"
    ]
    for col_num, val in enumerate(exemplo, 1):
        ws.cell(row=2, column=col_num, value=val)

    wb.save(caminho)


def importar_parcelas(caminho: str) -> tuple[list[Parcela], list[str]]:
    """
    Importa e valida parcelas de um arquivo Excel.
    Retorna uma tupla (lista_de_parcelas, lista_de_erros).
    """
    parcelas: list[Parcela] = []
    erros: list[str] = []

    if not os.path.exists(caminho):
        erros.append(f"Arquivo não encontrado: {caminho}")
        return parcelas, erros

    try:
        wb = openpyxl.load_workbook(caminho, data_only=True)
        ws = wb.active
    except Exception as e:
        erros.append(f"Erro ao abrir arquivo Excel: {e}")
        return parcelas, erros

    colunas = [cell.value for cell in ws[1]]
    linhas_lidas = []

    # Lê dados
    for row_idx in range(2, ws.max_row + 1):
        row_values = [ws.cell(row=row_idx, column=c).value for c in range(1, len(colunas) + 1)]
        if all(val is None for val in row_values):
            continue  # Pula linhas totalmente vazias
        row_dict = dict(zip(colunas, row_values))
        linhas_lidas.append((row_idx, row_dict))

    # Validações por linha
    numeros_vistos = set()
    for row_idx, row in linhas_lidas:
        linha_erros = []

        # 1. Validação de número e duplicado
        num = row.get("numero")
        if num is None:
            linha_erros.append("Número da parcela ausente.")
        else:
            try:
                num = int(num)
                if num in numeros_vistos:
                    linha_erros.append(f"Número de linha duplicada: {num}.")
                numeros_vistos.add(num)
            except ValueError:
                linha_erros.append("Número da parcela inválido.")

        # 2. Histórico ausente
        hist = row.get("historico")
        if not hist:
            linha_erros.append("Histórico ausente.")

        # 3. Valor bruto vazio ou negativo
        v_bruto = row.get("valor_bruto")
        if v_bruto is None:
            linha_erros.append("Valor bruto vazio.")
        else:
            try:
                v_bruto = Decimal(str(v_bruto))
                if v_bruto < 0:
                    linha_erros.append("Valor bruto negativo.")
            except Exception:
                linha_erros.append("Valor bruto inválido.")

        # 4. Valor pago maior que o valor bruto ou negativo
        v_pago = row.get("valor_pago_na_data")
        if v_pago is not None:
            try:
                v_pago = Decimal(str(v_pago))
                if v_pago < 0:
                    linha_erros.append("Valor pago na data negativo.")
                elif v_bruto is not None and isinstance(v_bruto, Decimal) and v_pago > v_bruto:
                    linha_erros.append("Valor pago na data maior que o valor bruto.")
            except Exception:
                linha_erros.append("Valor pago na data inválido.")
        else:
            v_pago = Decimal("0")

        # 5. Data de vencimento inválida
        dt_venc = row.get("data_vencimento")
        if dt_venc:
            if isinstance(dt_venc, datetime):
                dt_venc = dt_venc.date()
            elif isinstance(dt_venc, str):
                try:
                    dt_venc = datetime.strptime(dt_venc.strip(), "%Y-%m-%d").date()
                except ValueError:
                    try:
                        dt_venc = datetime.strptime(dt_venc.strip(), "%d/%m/%Y").date()
                    except ValueError:
                        linha_erros.append("Data de vencimento inválida (esperado YYYY-MM-DD ou DD/MM/AAAA).")
            elif not isinstance(dt_venc, date):
                linha_erros.append("Data de vencimento com tipo inválido.")
        else:
            linha_erros.append("Data de vencimento ausente.")

        # 6. Índice de correção inexistente
        corr = row.get("correcao_monetaria")
        if corr:
            corr_str = str(corr).strip().lower()
            try:
                # Tenta mapear ou achar correspondência
                IndiceCorrecao(corr_str)
            except ValueError:
                linha_erros.append(f"Índice de correção '{corr}' inexistente.")

        # 7. Tipo de juros inexistente
        jur = row.get("juros_moratorios")
        if jur:
            jur_str = str(jur).strip().lower()
            try:
                TipoJuros(jur_str)
            except ValueError:
                linha_erros.append(f"Tipo de juros '{jur}' inexistente.")

        # Data inicial correção e juros
        dt_corr = row.get("data_inicial_correcao")
        if dt_corr and isinstance(dt_corr, datetime):
            dt_corr = dt_corr.date()
        elif dt_corr and isinstance(dt_corr, str):
            try:
                dt_corr = datetime.strptime(dt_corr.strip(), "%Y-%m-%d").date()
            except ValueError:
                try:
                    dt_corr = datetime.strptime(dt_corr.strip(), "%d/%m/%Y").date()
                except ValueError:
                    pass

        dt_jur = row.get("data_inicial_juros")
        if dt_jur and isinstance(dt_jur, datetime):
            dt_jur = dt_jur.date()
        elif dt_jur and isinstance(dt_jur, str):
            try:
                dt_jur = datetime.strptime(dt_jur.strip(), "%Y-%m-%d").date()
            except ValueError:
                try:
                    dt_jur = datetime.strptime(dt_jur.strip(), "%d/%m/%Y").date()
                except ValueError:
                    pass

        if linha_erros:
            erros.append(f"Linha {row_idx}: " + "; ".join(linha_erros))
        else:
            parcela = Parcela(
                numero=int(num),
                natureza=str(row.get("natureza") or ""),
                data_vencimento=dt_venc,
                historico=str(hist),
                valor_bruto=Decimal(str(v_bruto)),
                valor_pago_na_data=Decimal(str(v_pago or 0)),
                correcao_monetaria=IndiceCorrecao(str(corr).strip().lower()) if corr else IndiceCorrecao.SEM_CORRECAO,
                data_inicial_correcao=dt_corr,
                juros_moratorios=TipoJuros(str(jur).strip().lower()) if jur else TipoJuros.SEM_JUROS,
                data_inicial_juros=dt_jur,
                multa_moratoria=Decimal(str(row.get("multa_moratoria") or 0)),
                observacao=str(row.get("observacao") or "")
            )
            parcela.recalcular_apurado()
            parcelas.append(parcela)

    return parcelas, erros
